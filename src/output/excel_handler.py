"""
Excel output handler for job scraper results.
"""

import os
import logging
from typing import Dict, List, Any, Optional, Union
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from src.utils.logger import setup_logger


class ExcelHandler:
    """
    Handler for saving job data to Excel files.
    """
    
    def __init__(self):
        """
        Initialize the Excel handler.
        """
        self.logger = setup_logger("output.excel")
    
    def save(self, results: Dict[str, List[Dict[str, Any]]], 
            filename: str = "job_listings.xlsx", **kwargs) -> str:
        """
        Save job data to Excel file.
        
        Args:
            results: Dictionary mapping scraper names to lists of jobs
            filename: Output filename
            **kwargs: Additional parameters for Excel output
            
        Returns:
            Path to the saved Excel file
        """
        try:
            # Check if we have any results
            if not results or all(not jobs for jobs in results.values()):
                self.logger.warning("No job data to save")
                return ""
            
            # Convert job data to DataFrames
            dfs = {}
            all_jobs = []
            
            for source, jobs in results.items():
                if not jobs:
                    continue
                
                df = pd.DataFrame(jobs)
                dfs[source] = df
                
                # Add all jobs to the combined DataFrame
                all_jobs.extend(jobs)
                
            if not all_jobs:
                self.logger.warning("No job data after processing")
                return ""
            
            # Create a combined DataFrame
            all_jobs_df = pd.DataFrame(all_jobs)
            
            # Define standard column order
            columns_order = [
                'job_id', 'title', 'company', 'location', 'experience', 
                'salary', 'job_type', 'posted_date', 'description_summary',
                'description', 'skills', 'industry', 'education', 'apply_url', 'source'
            ]
            
            # Function to reorder columns
            def reorder_columns(df):
                # Get columns that exist in the DataFrame
                available_columns = [col for col in columns_order if col in df.columns]
                
                # Add any columns not in our standard order
                extra_columns = [col for col in df.columns if col not in columns_order]
                available_columns.extend(extra_columns)
                
                # Return reordered DataFrame
                return df[available_columns]
            
            # Reorder columns in all DataFrames
            all_jobs_df = reorder_columns(all_jobs_df)
            
            for source in dfs:
                dfs[source] = reorder_columns(dfs[source])
            
            # Save to Excel with formatting
            with pd.ExcelWriter(filename, engine='openpyxl') as writer:
                # Write the All Jobs sheet first
                all_jobs_df.to_excel(writer, sheet_name='All Jobs', index=False)
                
                # Write individual scraper sheets
                for source, df in dfs.items():
                    df.to_excel(writer, sheet_name=source, index=False)
                
                # Apply formatting to all sheets
                for sheet_name in writer.sheets:
                    worksheet = writer.sheets[sheet_name]
                    
                    # Format header row
                    for cell in worksheet[1]:
                        cell.font = Font(bold=True, color="FFFFFF")
                        cell.fill = PatternFill(start_color="5B9BD5", end_color="5B9BD5", fill_type="solid")
                        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                        
                        # Add borders
                        cell.border = Border(
                            left=Side(style='thin'),
                            right=Side(style='thin'),
                            top=Side(style='thin'),
                            bottom=Side(style='thin')
                        )
                    
                    # Adjust column widths based on content
                    for idx, column in enumerate(all_jobs_df.columns if sheet_name == 'All Jobs' else dfs[sheet_name].columns):
                        col_letter = chr(65 + idx) if idx < 26 else chr(64 + idx//26) + chr(65 + idx%26)
                        
                        # Set width based on content type
                        if column == 'description':
                            worksheet.column_dimensions[col_letter].width = 100
                        elif column == 'description_summary':
                            worksheet.column_dimensions[col_letter].width = 70
                        elif column in ['title', 'company', 'skills']:
                            worksheet.column_dimensions[col_letter].width = 30
                        elif column in ['location', 'source', 'job_type']:
                            worksheet.column_dimensions[col_letter].width = 20
                        elif column in ['salary', 'experience', 'posted_date']:
                            worksheet.column_dimensions[col_letter].width = 15
                        else:
                            worksheet.column_dimensions[col_letter].width = 15
                            
                    # Freeze the header row
                    worksheet.freeze_panes = 'A2'
                    
                    # Add alternating row colors
                    for row_idx in range(2, worksheet.max_row + 1):
                        if row_idx % 2 == 0:  # Even rows
                            for col_idx in range(1, worksheet.max_column + 1):
                                cell = worksheet.cell(row=row_idx, column=col_idx)
                                cell.fill = PatternFill(start_color="EBF1FA", end_color="EBF1FA", fill_type="solid")
            
            self.logger.info(f"Successfully saved {len(all_jobs)} jobs to {filename}")
            return os.path.abspath(filename)
            
        except Exception as e:
            self.logger.error(f"Error saving to Excel: {e}")
            
            # Try to save as CSV if Excel fails
            try:
                csv_filename = filename.replace('.xlsx', '.csv')
                all_jobs_df.to_csv(csv_filename, index=False, encoding='utf-8')
                self.logger.info(f"Saved {len(all_jobs)} jobs to {csv_filename} (CSV fallback)")
                return os.path.abspath(csv_filename)
            except Exception as csv_err:
                self.logger.error(f"Error saving to CSV: {csv_err}")
                return ""